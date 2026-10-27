# -*- coding: utf-8 -*-
"""
ICON(icon.pcsline.co.kr) 3562 B/L Profitability Inquiry 실적 엑셀 내려받기.

사용법:
    python icon_actuals.py                 # 현재 시각 기준으로 오전/오후 자동 판정
    python icon_actuals.py --session AM    # 오전 이름으로 저장
    python icon_actuals.py --show          # 브라우저 창을 띄워서 동작 확인
    python icon_actuals.py --from 20260901 --to 20261231   # 기간 직접 지정
    python icon_actuals.py --export        # 수출 실적 (POL=KR, POD 비움) -> data/actuals_ex.json
    python icon_actuals.py --feeder        # 피더 실적 (전달·이번 달 따로) -> data/actuals_fd.json

조회 기간은 전달 1일 ~ 이번 달 말일이다. (출항 전 항차는 실적이 없으므로 다음 달은 보지 않는다)
결과: actuals/YYYY-MM-DD_AM.xlsx       원본 엑셀 (git 에 올리지 않음)
      data/actuals.json                항차·POL·POD 별 TEU·금액 합계 (뷰어용)

엑셀에 운임·수익·화주명이 들어 있으므로 원본은 actuals/ 에만 두고(.gitignore),
뷰어에는 항차·POL·POD 별 합계만 올린다.

    python icon_actuals.py --build-only    # 이미 받은 엑셀로 합계 파일만 다시 만들기
"""
import argparse
import calendar
import json
import re
import sys
import time
import traceback
import warnings
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

from scraper import load_config, now_kst, log, dump_debug, describe_page

BASE = Path(__file__).resolve().parent
OUT = BASE / "actuals"
SUM_OUT = BASE / "data" / "actuals.json"
SITE = "https://icon.pcsline.co.kr/"
MENU_NO = "3562"
POL_LIST = "CN,VN,JP,TH,HK"
# 수출(--export): POL 만 한국으로 바꾸고 POD 는 비운 채 조회한다.
POL_LIST_EX = "KR"
SUM_OUT_EX = BASE / "data" / "actuals_ex.json"
# 피더(--feeder): Type=FEEDER, 기준일 L.P.D, POL·POD 둘 다 CN,JP,KR,VN. 전달·이번 달을 한 달씩 조회해서
# 월 -> 수출/수입/삼국 -> 화주별로 합친다.
PORTS_FD = "CN,JP,KR,VN"
SUM_OUT_FD = BASE / "data" / "actuals_fd.json"
WAIT_MS = 20 * 60 * 1000   # 조회·다운로드는 길면 5분 넘게 걸린다고 해서 넉넉히

# --- 화면 element ---------------------------------------------------------
LOGIN_ID = "#mf_ibx_empCd"
LOGIN_PW = "#mf_sct_password"
LOGIN_GO = "#mf_btn_login"
# 메뉴 검색칸은 처음엔 숨어 있고, 위에 덮인 label 을 눌러야 나타난다.
MENU_LABEL = "#mf_wfm_header_ibx_menuNm_label"
MENU_INPUT = "#mf_wfm_header_ibx_menuNm_input"

P = "#mf_tac_layout_contents_00010640_body_"
DATE_FROM = P + "searchMapINPTM1_input"
DATE_TO = P + "searchMapINPTM2_input"
POL = P + "idpol"
POD = P + "idpod"
# 조회조건의 Y/N/ALL 선택칸. 화면 기본값은 Y 인데 ALL 로 바꿔서 조회한다.
CLO = P[1:] + "searchMapINPCLO"
CLO_TEXT = "ALL"
# 피더 조회용. TYP 는 여러 개를 고르는 checkcombobox (FEEDER 의 값은 'F'), CLS 는 O.B.D / L.P.D 선택칸.
TYP = P[1:] + "searchMapINPTYP"
TYP_TEXT = "FEEDER"
CLS = P[1:] + "searchMapINPCLS"
CLS_TEXT = "L.P.D"
INQUIRY = P + "btn_search"
TOTAL_CNT = P + "totalCnt"
# 다운로드 버튼 id 가운데(wq_uuid_NNNN)는 화면을 열 때마다 바뀐다.
EXCEL = "a[id^='mf_tac_layout_contents_00010640_body_'][id$='_btnExcelDown']:visible"
ALERT_MSG = "[id^='mf_tac_layout_contents_00010640_body_alert'][id$='_tbx_message']"

# Inquiry 가 실제로 쏘는 조회 요청. 화면의 'Progress...' 창(processMsgLayer)은
# 자동 실행 브라우저에서는 DOM 에 잡히지 않아서, 대신 이 응답과 건수 표시를 기다린다.
QUERY_API = "/css/doc/selectCSSDOC3562R02_NEW"


def date_range(today):
    """2026-10-15 -> ('20260901', '20261031')  전달 1일 ~ 이번 달 말일"""
    y, m = today.year, today.month
    py, pm = (y - 1, 12) if m == 1 else (y, m - 1)
    return "%d%02d01" % (py, pm), "%d%02d%02d" % (y, m, calendar.monthrange(y, m)[1])


def check_alert(page):
    """화면 안 Alert 창 문구가 있으면 돌려준다."""
    for el in page.locator(ALERT_MSG).all():
        if el.is_visible():
            return el.inner_text().strip()
    return None


def total_count(page):
    txt = page.locator(TOTAL_CNT).inner_text() if page.locator(TOTAL_CNT).count() else ""
    digits = re.sub(r"[^0-9]", "", txt)
    return int(digits) if digits else 0


def type_into(page, sel, text):
    page.click(sel)
    page.keyboard.press("Control+A")
    page.keyboard.type(text, delay=60)
    page.keyboard.press("Tab")
    page.wait_for_timeout(500)


def select_text(page, cid, text):
    """WebSquare 선택칸(w2selectbox)을 보이는 글자로 고른다. 고른 뒤의 글자를 돌려준다."""
    return page.evaluate("""([cid, text]) => {
        const c = $p.getComponentById(cid);
        for (let i = 0; i < c.getItemCount(); i++)
            if (c.getItemText(i) === text) { c.setSelectedIndex(i); break; }
        return c.getText();
    }""", [cid, text])


def login(page, cfg):
    acc = cfg.get("icon_login") or {}
    uid, pwd = (acc.get("user_id") or "").strip(), (acc.get("password") or "").strip()
    if not uid or not pwd:
        raise RuntimeError("config.json 에 icon_login 아이디/비밀번호가 없습니다.")
    page.goto(SITE, wait_until="domcontentloaded", timeout=90000)
    page.wait_for_selector(LOGIN_ID, state="visible", timeout=60000)
    page.fill(LOGIN_ID, uid)
    page.fill(LOGIN_PW, pwd)
    page.click(LOGIN_GO)
    try:
        page.wait_for_selector(MENU_LABEL, state="visible", timeout=60000)
    except Exception:
        raise RuntimeError("로그인 실패 - 아이디/비밀번호를 확인하세요")
    page.wait_for_timeout(5000)   # 첫 화면 위젯들이 다 뜰 때까지
    log("ICON 로그인 성공 (%s)" % uid)


def open_menu(page):
    page.click(MENU_LABEL)
    page.wait_for_selector(MENU_INPUT, state="visible", timeout=10000)
    page.keyboard.type(MENU_NO, delay=120)
    page.wait_for_timeout(1500)
    page.keyboard.press("Enter")
    page.wait_for_selector(INQUIRY, state="visible", timeout=60000)
    page.wait_for_timeout(3000)
    log("  %s B/L Profitability Inquiry 화면 열림" % MENU_NO)


def inquiry(page, errors, allow_zero=False):
    """Inquiry 를 누르고 조회 결과 건수를 돌려준다. 서버 오류면 예외.

    allow_zero: 0건이어도 실패로 보지 않는다 (피더 당월 조회는 월초에 0건일 수 있다).
    """
    errors.clear()
    t0 = time.time()
    with page.expect_response(lambda r: r.url.endswith(QUERY_API), timeout=WAIT_MS) as info:
        page.click(INQUIRY)
    if info.value.status != 200:
        raise RuntimeError("조회 응답 HTTP %s" % info.value.status)
    # 응답에 든 행 수. 같은 화면에서 연달아 조회하면 화면 건수가 앞 조회 값으로 남아 있어서
    # 응답 행 수와 같아질 때까지 기다린다. 응답을 못 읽으면 예전처럼 0 이 아닌 건수를 기다린다.
    try:
        want = len(info.value.json().get("dlt_resultList") or [])
    except Exception:
        want = None
    if want == 0:
        if allow_zero:
            log("  Inquiry 완료: 0건 (%.0f초)" % (time.time() - t0))
            return 0
        raise RuntimeError("조회 결과가 0건입니다")
    # 응답을 받은 뒤에도 화면이 표를 채우고 보조 조회를 몇 번 더 한 다음에 건수를 바꾼다.
    for _ in range(180):
        if errors:
            raise RuntimeError("서버 오류: %s" % errors[-1])
        msg = check_alert(page)
        if msg:
            raise RuntimeError("조회 경고: %s" % msg)
        n = total_count(page)
        if n and (want is None or n == want):
            log("  Inquiry 완료: %d건 (%.0f초)" % (n, time.time() - t0))
            return n
        page.wait_for_timeout(1000)
    raise RuntimeError("조회 결과가 0건입니다" if want is None else
                       "화면 건수가 응답 건수(%d)로 바뀌지 않았습니다" % want)


def read_rows(path):
    """엑셀을 (머리글 4행 이름 -> 열 번호, 6행부터의 데이터 행) 으로 읽는다."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")   # 사이트 엑셀에 기본 스타일이 없다는 경고
        ws = openpyxl.load_workbook(path).active
    head = {}
    for i, c in enumerate(ws[4]):
        if c.value and c.value not in head:
            head[c.value] = i
    return head, [r for r in ws.iter_rows(min_row=6, values_only=True) if any(r)]


def summarize(path, export=False):
    """항차(VVD) -> POL -> POD 별 TEU(U열) 합계와 Income S.TTL(AD열) 합계.

    POD 가 한국(KR) 이면 항구마다 따로, 그 밖은 나라별로 합친다 (RUVVO+RUVFP -> RU).
    RPT(= 금액 합계 / TEU 합계)는 뷰어에서 나눠 보여준다.
    Type(Q열) 이 FEE 인 행은 어떤 경우에도 넣지 않는다.
    """
    head, rows = read_rows(path)
    c_vvd, c_pol, c_pod = head.get("VVD", 6), head.get("POL", 13), head.get("POD", 14)
    c_teu = head.get("TEU", 20)
    c_type = head.get("Type", 16)
    c_obd = head.get("On Board\nDate", 10)
    c_amt = 29   # AD열 = Income 아래 S.TTL (5행 머리글이라 4행 이름으로는 못 찾는다)
    num = lambda v: float(v) if isinstance(v, (int, float)) else 0.0

    agg, first_obd = {}, {}
    for r in rows:
        vvd = (r[c_vvd] or "").strip()
        pol = (r[c_pol] or "").strip()
        pod = (r[c_pod] or "").strip()
        if not vvd or not pol:
            continue
        if str(r[c_type] or "").strip().upper() == "FEE":
            continue
        # 수출은 POD 가 외국 항구라 나라로 합치지 않고 항구별로 둔다
        key = pod if export or pod.startswith("KR") else pod[:2]
        g = agg.setdefault(vvd, {}).setdefault((pol, key), {"teu": 0.0, "amt": 0.0, "bl": 0, "pods": set()})
        g["teu"] += num(r[c_teu])
        g["amt"] += num(r[c_amt])
        g["bl"] += 1
        g["pods"].add(pod)
        # POL 순서를 정할 때 스케줄에 없는 POL 은 가장 이른 On Board Date 로 대신한다
        obd = str(r[c_obd] or "")[:10]
        if obd and (vvd, pol) not in first_obd or obd and obd < first_obd[(vvd, pol)]:
            first_obd[(vvd, pol)] = obd

    out = {}
    for vvd, groups in agg.items():
        # POL 순, 같은 POL 안에서는 한국 항구 먼저, 그다음 나라 코드 순
        lines = sorted(groups.items(), key=lambda kv: (kv[0][0], not kv[0][1].startswith("KR"), kv[0][1]))
        out[vvd] = [{
            "pol": pol, "pod": key, "pods": sorted(g["pods"]),
            "teu": round(g["teu"], 2), "amt": round(g["amt"], 2), "bl": g["bl"],
            "obd": first_obd.get((vvd, pol), ""),
        } for (pol, key), g in lines]
    return out


def build_summary(xlsx, period, export=False):
    """뷰어용 항차별 합계 파일을 만든다. (비밀번호는 나중에 붙이기로 해서 지금은 평문)"""
    summary = summarize(xlsx, export)
    sum_out = SUM_OUT_EX if export else SUM_OUT
    payload = {
        "updated": now_kst().strftime("%Y-%m-%d %H:%M"),
        "period": period,
        "source": xlsx.name,
        "vvd": summary,
    }
    sum_out.parent.mkdir(exist_ok=True)
    sum_out.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    log("  뷰어용 실적 저장: data/%s (항차 %d개)" % (sum_out.name, len(summary)))
    return True


def latest_xlsx(export=False):
    files = sorted(OUT.glob("????-??-??_[AP]M%s.xlsx" % ("_EX" if export else "")))
    return files[-1] if files else None


# --- 피더 -------------------------------------------------------------------
# 화주 = 엑셀 AO열(Shipper Name), 없으면 AQ열(Consignee Name) 이름으로 가린다.
# (AN·AP 열은 각각의 거래처 코드라 이름으로 판정한다)
# 한 이름 안에 둘이 같이 나오면 앞에 있는 것이 먼저다.
#   'APL CO PTE (CNC) C/O CMA CGM KOREA' -> APL
# 단, Shipper 가 APL·ANL 이고 Consignee 가 CMA 인 B/L 은 CMA 로 본다 (사용자 지정).
CARRIERS = [
    ("APL", re.compile(r"\bAPL\b")),
    ("ANL", re.compile(r"\bANL\b")),
    ("CMA", re.compile(r"\bCMA\b")),
    ("MSC", re.compile(r"\bM\.?S\.?C\b")),
    ("ZIM", re.compile(r"\bZIM\b")),
    ("HMM", re.compile(r"\bHMM\b|HYUNDAI MERCHANT")),
    ("HUENG A LINE", re.compile(r"\bH(?:EU|UE)NG[\s-]?A\b")),
    ("YANGMING", re.compile(r"YANG\s?MING")),
    ("DONGJIN", re.compile(r"\bDONG\s?JIN\b")),
    ("SOFAST KOREA", re.compile(r"\bSOFAST\b")),
    ("UNITED LINE", re.compile(r"\bUNITED LINE\b")),
    ("S1 SHIPPING", re.compile(r"\bS1 SHIPPING\b")),
    ("PAN CONTINENTAL", re.compile(r"\bPAN\s?CON(?:TINENTAL)?\b")),
]
CARRIER_ETC = "기타"
# 뷰어 표의 화주 순서 (사용자가 불러준 순서)
CARRIER_ORDER = ["MSC", "CMA", "APL", "ANL", "ZIM", "HMM", "HUENG A LINE", "YANGMING",
                 "DONGJIN", "SOFAST KOREA", "UNITED LINE", "S1 SHIPPING", "PAN CONTINENTAL", CARRIER_ETC]
SKIP_RE = re.compile(r"NAM\s?SUNG")   # 남성해운 건은 집계하지 않는다
C_SHIPPER, C_CONSIGNEE = 40, 42      # AO, AQ


def carrier_of(*names):
    for name in names:
        name = (name or "").upper()
        for label, rx in CARRIERS:
            if rx.search(name):
                return label
    return None


def month_ranges(today):
    """2026-10-10 -> [('2026-09', '20260901', '20260930'), ('2026-10', '20261001', '20261031')]"""
    y, m = today.year, today.month
    py, pm = (y - 1, 12) if m == 1 else (y, m - 1)
    return [("%d-%02d" % (yy, mm), "%d%02d01" % (yy, mm), "%d%02d%02d" % (yy, mm, calendar.monthrange(yy, mm)[1]))
            for yy, mm in ((py, pm), (y, m))]


def feeder_dir(pol, pod):
    """수출 = POL 이 KR, 수입 = POD 가 KR, 삼국 = 둘 다 KR 이 아님. KR -> KR(연안) 은 None (집계 안 함)"""
    if pol.startswith("KR") and pod.startswith("KR"):
        return None
    if pol.startswith("KR"):
        return "EX"
    if pod.startswith("KR"):
        return "IM"
    return "TS"


def summarize_feeder(path):
    """방향(EX/IM/TS) -> 화주 -> TEU(U열)·Income S.TTL(AD열) 합계, 그리고 분류 못 한 이름 목록.

    피더 B/L 은 Type(Q열) 이 전부 FEE 라서 수입·수출 실적과 달리 FEE 를 빼지 않는다.
    """
    head, rows = read_rows(path)
    c_pol, c_pod, c_teu = head.get("POL", 13), head.get("POD", 14), head.get("TEU", 20)
    num = lambda v: float(v) if isinstance(v, (int, float)) else 0.0
    dirs, etc, skipped = {}, {}, 0
    for r in rows:
        pol = (r[c_pol] or "").strip().upper()
        pod = (r[c_pod] or "").strip().upper()
        if not pol or not pod:
            continue
        shp, cne = (r[C_SHIPPER] or "").strip(), (r[C_CONSIGNEE] or "").strip()
        if SKIP_RE.search(shp.upper()) or SKIP_RE.search(cne.upper()):
            skipped += 1
            continue
        d = feeder_dir(pol, pod)
        if not d:
            continue
        who = carrier_of(shp, cne)
        if who in ("APL", "ANL") and carrier_of(cne) == "CMA":
            who = "CMA"
        if not who:
            who = CARRIER_ETC
            e = etc.setdefault((d, shp, cne), {"dir": d, "shipper": shp, "consignee": cne, "teu": 0.0, "bl": 0})
            e["teu"] += num(r[c_teu])
            e["bl"] += 1
        g = dirs.setdefault(d, {}).setdefault(who, {"teu": 0.0, "amt": 0.0, "bl": 0})
        g["teu"] += num(r[c_teu])
        g["amt"] += num(r[29])   # AD열 = Income S.TTL
        g["bl"] += 1
    for groups in dirs.values():
        for g in groups.values():
            g["teu"], g["amt"] = round(g["teu"], 2), round(g["amt"], 2)
    return dirs, sorted(etc.values(), key=lambda e: -e["bl"]), skipped


def build_feeder(files):
    """files: {'2026-09': (xlsx, ['20260901','20260930']), ...}

    이번에 조회한 달만 새로 만들고, 예전에 받아 둔 달은 기존 파일에서 그대로 둔다.
    """
    old = {}
    if SUM_OUT_FD.exists():
        try:
            old = json.loads(SUM_OUT_FD.read_text(encoding="utf-8")).get("months") or {}
        except Exception:
            old = {}
    months = dict(old)
    for ym, (xlsx, period) in files.items():
        # 조회 결과가 0건이면 받은 엑셀이 없다 (xlsx=None)
        dirs, etc, skipped = summarize_feeder(xlsx) if xlsx else ({}, [], 0)
        months[ym] = {"period": period, "source": xlsx.name if xlsx else "", "dirs": dirs, "etc": etc}
        log("  피더 %s: %s" % (ym, ", ".join("%s %d건" % (d, sum(g["bl"] for g in gs.values()))
                                             for d, gs in sorted(dirs.items())) or "0건")
            + ("  (NAM SUNG %d건 제외)" % skipped if skipped else ""))
        for e in etc:
            log("    분류 못 한 화주(%s): %s / %s (%d건, %s TEU)" % (e["dir"], e["shipper"], e["consignee"], e["bl"], e["teu"]))
    payload = {
        "updated": now_kst().strftime("%Y-%m-%d %H:%M"),
        "carriers": CARRIER_ORDER,
        "months": dict(sorted(months.items())),
    }
    SUM_OUT_FD.parent.mkdir(exist_ok=True)
    SUM_OUT_FD.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    log("  뷰어용 피더 실적 저장: data/%s (%d개월)" % (SUM_OUT_FD.name, len(months)))
    return True


def latest_feeder_files():
    """달마다 가장 최근에 받은 피더 엑셀 (--build-only 용)"""
    out = {}
    for f in sorted(OUT.glob("????-??-??_[AP]M_FD_??????.xlsx")):
        ym = f.stem[-6:-2] + "-" + f.stem[-2:]
        out[ym] = (f, None)
    return out


def check_only(page, cid, text):
    """WebSquare checkcombobox 에서 보이는 글자가 text 인 항목 하나만 체크한다. 고른 뒤의 글자를 돌려준다."""
    return page.evaluate("""([cid, text]) => {
        const c = $p.getComponentById(cid);
        const it = c.getItemArray().find(x => x.label === text);
        if (it) c.setValue(it.value);
        return c.getText();
    }""", [cid, text])


def open_browser(pw, cfg, show):
    """ICON 화면용 브라우저. (browser, page, errors) 를 돌려준다."""
    headless = cfg.get("headless", True) and not show
    browser = pw.chromium.launch(headless=headless)
    # 화면이 좁으면 조회조건 줄이 가로로 밀려 엉뚱한 칸이 눌린다
    ctx = browser.new_context(viewport={"width": 2300, "height": 1000},
                              locale="ko-KR", accept_downloads=True)
    page = ctx.new_page()
    # 사이트가 서버 오류를 브라우저 경고창으로 띄운다 ("서버에러입니다(timeout)").
    # 첫 화면 위젯에서도 뜨므로 조회 직전에 비우고, 조회 중에 뜬 것만 실패로 본다.
    errors = []

    def on_dialog(d):
        m = d.message.strip()
        log("  브라우저 경고: %s" % m)
        if "에러" in m or "error" in m.lower():
            errors.append(m)
        d.accept()
    page.on("dialog", on_dialog)
    return browser, page, errors


def download(page, errors, tmp, n):
    """Excel Download 를 눌러 tmp 에 저장하고 엑셀 건수를 검증한다."""
    t0 = time.time()
    errors.clear()
    with page.expect_download(timeout=WAIT_MS) as dl:
        page.locator(EXCEL).first.click()
    dl.value.save_as(str(tmp))
    got = count_rows(tmp)
    log("  엑셀 다운로드 완료: %d건 (%.0f초)" % (got, time.time() - t0))
    if got == 0:
        raise RuntimeError("내려받은 엑셀이 비어 있습니다")
    if got != n:
        log("  경고: 화면 건수(%d)와 엑셀 건수(%d)가 다릅니다" % (n, got))


def inquiry_retry(page, errors, allow_zero=False):
    # 서버가 가끔 timeout 을 낸다. 한 번 더 눌러 본다.
    for attempt in (1, 2):
        try:
            return inquiry(page, errors, allow_zero)
        except Exception as exc:
            log("  조회 실패 (%d/2): %s" % (attempt, exc))
            if attempt == 2:
                raise
            page.wait_for_timeout(10000)


def run_feeder(session=None, show=False):
    """피더 실적: 전달·이번 달을 한 달씩 따로 조회해서 받는다."""
    cfg = load_config()
    now = now_kst()
    if session is None:
        session = "AM" if now.hour < 12 else "PM"
    session = session.upper()
    months = month_ranges(now.date())

    log("=" * 60)
    log("ICON 피더 실적 다운로드 시작: %s %s  (%s, POL/POD %s, %s)"
        % (now.strftime("%Y.%m.%d"), session, " / ".join(m for m, _, _ in months), PORTS_FD, CLS_TEXT))
    OUT.mkdir(exist_ok=True)
    got = {}

    with sync_playwright() as pw:
        browser, page, errors = open_browser(pw, cfg, show)
        tmp = None
        try:
            login(page, cfg)
            open_menu(page)

            typ = check_only(page, TYP, TYP_TEXT)
            if typ != TYP_TEXT:
                raise RuntimeError("INPTYP 를 %s 만 체크하지 못했습니다 (현재 %s)" % (TYP_TEXT, typ))
            cls = select_text(page, CLS, CLS_TEXT)
            if cls != CLS_TEXT:
                raise RuntimeError("INPCLS 선택칸을 %s 로 바꾸지 못했습니다 (현재 %s)" % (CLS_TEXT, cls))
            clo = select_text(page, CLO, CLO_TEXT)
            if clo != CLO_TEXT:
                raise RuntimeError("INPCLO 선택칸을 %s 로 바꾸지 못했습니다 (현재 %s)" % (CLO_TEXT, clo))
            type_into(page, POL, PORTS_FD)
            type_into(page, POD, PORTS_FD)

            for ym, d_from, d_to in months:
                type_into(page, DATE_FROM, d_from)
                type_into(page, DATE_TO, d_to)
                log("  조회조건: %s ~ %s, POL=%s, POD=%s, INPTYP=%s, INPCLS=%s, INPCLO=%s"
                    % (page.input_value(DATE_FROM), page.input_value(DATE_TO), page.input_value(POL),
                       page.input_value(POD), typ, cls, clo))
                n = inquiry_retry(page, errors, allow_zero=True)
                out = OUT / ("%s_%s_FD_%s.xlsx" % (now.strftime("%Y-%m-%d"), session, ym.replace("-", "")))
                if n == 0:
                    # 엑셀을 받을 게 없다. 그달 실적은 비어 있는 것으로 저장한다.
                    got[ym] = (None, [d_from, d_to])
                    continue
                tmp = out.with_suffix(".part.xlsx")
                download(page, errors, tmp, n)
                tmp.replace(out)
                tmp = None
                log("저장 완료: actuals/%s (%.0f KB)" % (out.name, out.stat().st_size / 1024))
                got[ym] = (out, [d_from, d_to])
        except Exception:
            log("  현재 화면: %s" % describe_page(page))
            dump_debug(page, "icon_feeder_fatal")
            if tmp:
                tmp.unlink(missing_ok=True)
            raise
        finally:
            browser.close()

    build_feeder(got)
    return 0


def count_rows(path):
    """내려받은 엑셀의 B/L 건수. 머리글이 4~5행이고 6행부터 데이터다."""
    head, rows = read_rows(path)
    col = head.get("B/L No.", 7)
    return sum(1 for r in rows if r[col])


def run(session=None, show=False, d_from=None, d_to=None, export=False):
    cfg = load_config()
    now = now_kst()
    if session is None:
        session = "AM" if now.hour < 12 else "PM"
    session = session.upper()
    auto_from, auto_to = date_range(now.date())
    d_from, d_to = d_from or auto_from, d_to or auto_to
    pol_list = POL_LIST_EX if export else POL_LIST

    log("=" * 60)
    log("ICON %s 실적 다운로드 시작: %s %s  (기간 %s ~ %s, POL %s)"
        % ("수출" if export else "수입", now.strftime("%Y.%m.%d"), session, d_from, d_to, pol_list))

    OUT.mkdir(exist_ok=True)
    out = OUT / ("%s_%s%s.xlsx" % (now.strftime("%Y-%m-%d"), session, "_EX" if export else ""))
    tmp = out.with_suffix(".part.xlsx")

    with sync_playwright() as pw:
        browser, page, errors = open_browser(pw, cfg, show)
        try:
            login(page, cfg)
            open_menu(page)

            type_into(page, DATE_FROM, d_from)
            type_into(page, DATE_TO, d_to)
            type_into(page, POL, pol_list)
            page.fill(POD, "")
            clo = select_text(page, CLO, CLO_TEXT)
            if clo != CLO_TEXT:
                raise RuntimeError("INPCLO 선택칸을 %s 로 바꾸지 못했습니다 (현재 %s)" % (CLO_TEXT, clo))
            log("  조회조건: %s ~ %s, POL=%s, POD=(비움), INPCLO=%s"
                % (page.input_value(DATE_FROM), page.input_value(DATE_TO), page.input_value(POL), clo))

            n = inquiry_retry(page, errors)
            download(page, errors, tmp, n)
        except Exception:
            log("  현재 화면: %s" % describe_page(page))
            dump_debug(page, "icon_fatal")
            tmp.unlink(missing_ok=True)
            raise
        finally:
            browser.close()

    # 검증을 통과한 파일만 최종 이름으로 바꾼다. 실패하면 이전 파일이 그대로 남는다.
    tmp.replace(out)
    log("저장 완료: actuals/%s (%.0f KB)" % (out.name, out.stat().st_size / 1024))
    build_summary(out, [d_from, d_to], export)
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="ICON B/L Profitability 실적 엑셀 다운로드")
    ap.add_argument("--session", choices=["AM", "PM", "am", "pm"], default=None,
                    help="오전(AM)/오후(PM) 이름으로 저장. 생략하면 현재 시각으로 판정")
    ap.add_argument("--show", action="store_true", help="브라우저 창을 띄워서 실행")
    ap.add_argument("--from", dest="d_from", help="조회 시작일 YYYYMMDD (생략하면 스케줄 조회 월 기준)")
    ap.add_argument("--to", dest="d_to", help="조회 종료일 YYYYMMDD")
    ap.add_argument("--build-only", action="store_true",
                    help="조회하지 않고 가장 최근 엑셀로 뷰어용 합계 파일만 다시 만든다")
    ap.add_argument("--export", action="store_true",
                    help="수출 실적: POL=KR, POD 비움. actuals/*_EX.xlsx, data/actuals_ex.json 에 따로 저장")
    ap.add_argument("--feeder", action="store_true",
                    help="피더 실적: FEEDER, L.P.D, POL/POD=CN,JP,KR,VN 로 전달·이번 달을 따로 조회 -> data/actuals_fd.json")
    args = ap.parse_args()
    if args.feeder:
        if args.build_only:
            files = latest_feeder_files()
            if not files:
                raise SystemExit("actuals/ 에 받은 피더 엑셀이 없습니다.")
            sys.exit(0 if build_feeder(files) else 1)
        try:
            sys.exit(run_feeder(args.session, args.show))
        except Exception as e:
            log("ICON 피더 실적 치명적 오류: %s" % e)
            traceback.print_exc()
            sys.exit(2)
    if args.build_only:
        x = latest_xlsx(args.export)
        if not x:
            raise SystemExit("actuals/ 에 받은 엑셀이 없습니다.")
        sys.exit(0 if build_summary(x, None, args.export) else 1)
    try:
        sys.exit(run(args.session, args.show, args.d_from, args.d_to, args.export))
    except Exception as e:
        log("ICON 실적 치명적 오류: %s" % e)
        traceback.print_exc()
        sys.exit(2)
